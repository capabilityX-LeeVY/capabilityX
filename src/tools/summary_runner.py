import os
import traceback
import pandas as pd
import sqlite3
import platform

from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from config.mysql import ConnectToDB
from common.utils.logger import LOGGER_INFO
from config.paths import CONTROL_TYPE_PATH, DB_PATH


class ControlTypeExporter:
    def __init__(self):
        self.folder_path = CONTROL_TYPE_PATH
        if not os.path.exists(self.folder_path):
            os.makedirs(self.folder_path)
        self.win_db_path = DB_PATH
        if not os.path.exists(self.win_db_path) and platform.system().lower() == 'windows':
            raise FileNotFoundError(f"Windows DB not found: {self.win_db_path}")

    def _save_to_excel_with_chart(self, df, filename, chart_title, category_col, value_col):
        file_path = os.path.join(self.folder_path, filename)
        try:
            df.to_excel(file_path, index=False)
            wb = load_workbook(file_path)
            ws = wb.active
            if ws is not None:
                max_row = ws.max_row
            else:
                LOGGER_INFO("Worksheet is not initialized properly.")
                return
            data_max_row = max_row - 1
            chart = BarChart()
            chart.title = chart_title
            chart.x_axis.title = category_col
            chart.y_axis.title = value_col

            value_col_index = df.columns.get_loc(value_col) + 1
            category_col_index = df.columns.get_loc(category_col) + 1

            data = Reference(ws, min_col=value_col_index, min_row=1, max_row=data_max_row)
            categories = Reference(ws, min_col=category_col_index, min_row=2, max_row=data_max_row)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(categories)

            # chart.legend = None
            chart.style = 10
            chart.height = 10
            chart.width = 20
            chart.dataLabels = DataLabelList()
            chart.dataLabels.showVal = True
            chart.dataLabels.position = "outEnd"
            chart.dataLabels.showSerName = False
            chart.dataLabels.showCatName = False
            chart.x_axis.tickLblPos = "low"

            if ws is not None:
                ws.add_chart(chart, "E2") #type: ignore
            else:
                LOGGER_INFO("Worksheet is not initialized properly.")
            wb.save(file_path)

            LOGGER_INFO(f"Excel with chart saved: {file_path}")
        except Exception as e:
            LOGGER_INFO(f"Error saving Excel with chart: {e}")
            LOGGER_INFO(traceback.format_exc())

    def _query_mac_db(self):
        sql = '''
        SELECT
            Role as "Control Type",
            COUNT(*) AS "Total Controls",
            COUNT(DISTINCT controlID) AS "Distinct Controls"
        FROM navigation_mac_deep_baseline
        GROUP BY Role
        ORDER BY 2 DESC
        '''
        conn = ConnectToDB()
        try:
            results, columns = conn.Query(sql)
            return pd.DataFrame(results, columns=columns)
        except Exception as e:
            LOGGER_INFO(f"Error querying mac DB: {e}")
            return pd.DataFrame()

    def _query_win_db(self):
        sql = '''
        SELECT "ControlLocalizedType" as "Control Type", 
               COUNT(*) as "Total Controls", 
               COUNT(DISTINCT "Description") as "Distinct Controls"
        FROM "ControlINFO" 
        GROUP BY "ControlLocalizedType" 
        ORDER BY "Total Controls" DESC
        '''
        try:
            conn = sqlite3.connect(self.win_db_path)
            df = pd.read_sql_query(sql, conn)
            conn.close()
            return df
        except Exception as e:
            LOGGER_INFO(f"Error querying win DB: {e}")
            return pd.DataFrame()

    def export_mac_controltype_summary(self):
        df = self._query_mac_db()
        if df.empty:
            LOGGER_INFO("No mac control type data found")
            return
        total_count = df['Total Controls'].sum()
        distinct_count = df['Distinct Controls'].sum()
        total_row = pd.DataFrame([['Total', total_count, distinct_count]], columns=df.columns)
        df = pd.concat([df, total_row], ignore_index=True)
        self._save_to_excel_with_chart(
            df,
            "macOS_data_TotalforEachControltype.xlsx",
            "Total Count by Control Type (Mac)",
            category_col='Control Type',
            value_col='Total Controls'
        )

    def export_win_controltype_summary(self):
        df = self._query_win_db()
        if df.empty:
            LOGGER_INFO("No win control type data found")
            return
        total_count = df['Total Controls'].sum()
        distinct_count = df['Distinct Controls'].sum()
        total_row = pd.DataFrame([['Total', total_count, distinct_count]], columns=df.columns)
        df = pd.concat([df, total_row], ignore_index=True)
        self._save_to_excel_with_chart(
            df,
            "windowsOS_data_TotalforEachControlType.xlsx",
            "Total Count by Control Type (Windows)",
            category_col='Control Type',
            value_col='Total Controls'
        )


if __name__ == '__main__':
    exporter = ControlTypeExporter()
    exporter.export_mac_controltype_summary()
    exporter.export_win_controltype_summary()
